"""Exports Excel des instruments, aux formats constatés dans les ministères.

Le PPM reprend la présentation officielle à deux onglets (fournitures, travaux et
services ; prestations intellectuelles) avec une colonne de date par étape.
"""

from __future__ import annotations

from collections import Counter
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.engines.instruments import build_pta, consumption_plan
from app.engines.procurement import MarketType, ProcurementLot
from app.engines.snapshot import PlanningSnapshot

HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
BOLD = Font(bold=True)
MONTHS = ["Janv", "Févr", "Mars", "Avr", "Mai", "Juin", "Juil", "Août", "Sept", "Oct", "Nov",
          "Déc"]  # fmt: skip


def _header(ws, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for col in range(1, len(labels) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.column_dimensions["B"].width = 50


def _save(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_pta(snapshot: PlanningSnapshot, title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "PTA"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    _header(ws, 3, ["Action", "Activité", "Tâche", "Ressource", "Quantité", "Unité",
                    "Prix unitaire (FCFA)", "Montant (FCFA)", "Imputation", "Source",
                    "Mode d'exécution", "Trimestre"])  # fmt: skip
    for i, row in enumerate(build_pta(snapshot), start=4):
        values = [row.action_code, row.activity_label, row.task_label, row.resource_label,
                  float(row.quantity), row.unit, row.unit_price, row.cost, row.budget_line,
                  row.funding_source, row.execution_mode, f"T{row.quarter}"]  # fmt: skip
        for col, value in enumerate(values, start=1):
            ws.cell(row=i, column=col, value=value)
    return _save(wb)


def export_pcc(snapshot: PlanningSnapshot, title: str) -> bytes:
    wb = Workbook()
    monthly = wb.active
    monthly.title = "PCC mensuel"
    monthly["A1"] = title
    monthly["A1"].font = Font(bold=True, size=13)
    _header(monthly, 3, ["Imputation", *MONTHS, "Total"])
    quarterly = wb.create_sheet("PCC trimestriel")
    _header(quarterly, 1, ["Imputation", "T1", "T2", "T3", "T4", "Total"])
    for i, (line, months) in enumerate(consumption_plan(snapshot).items()):
        monthly.cell(row=4 + i, column=1, value=line)
        for m, amount in enumerate(months, start=2):
            monthly.cell(row=4 + i, column=m, value=amount)
        monthly.cell(row=4 + i, column=14, value=sum(months))
        quarters = [sum(months[q * 3 : q * 3 + 3]) for q in range(4)]
        for q, amount in enumerate([line, *quarters, sum(months)], start=1):
            quarterly.cell(row=2 + i, column=q, value=amount)
    return _save(wb)


PPM_BASE = ["N°", "Réf N°", "Description", "Type de marché", "Mode de passation",
            "Montant estimatif HT (FCFA)", "Source de financement", "Ligne d'imputation",
            "Organe de contrôle", "Autorisation d'engagement"]  # fmt: skip


def export_ppm(snapshot: PlanningSnapshot, lots: list[ProcurementLot], authority: str) -> bytes:
    lines = {line.id: line for line in snapshot.lines}
    activities = {a.id: a.label for a in snapshot.activities}
    wb = Workbook()
    fts = wb.active
    fts.title = "Fournitures, travaux, services"
    pi = wb.create_sheet("Prestations intellectuelles")
    counters = {fts.title: 0, pi.title: 0}
    for ws in (fts, pi):
        ws["A1"] = f"Nom de l'Autorité Contractante : {authority}"
        ws["A2"] = f"Période couverte : 1er janvier au 31 décembre {snapshot.fiscal_year}"
        ws["A1"].font = BOLD

    ordered = sorted(lots, key=lambda lot: (lot.launch_date, lot.id))
    for ws, keep in ((fts, lambda t: t != MarketType.PRESTATIONS_INTELLECTUELLES),
                     (pi, lambda t: t == MarketType.PRESTATIONS_INTELLECTUELLES)):  # fmt: skip
        selected = [lot for lot in ordered if keep(lot.market_type)]
        step_labels = []
        for lot in selected:
            for step in lot.steps:
                if step.label not in step_labels:
                    step_labels.append(step.label)
        _header(ws, 4, PPM_BASE + [f"Date — {label}" for label in step_labels]
                + ["Activités PTA d'origine"])  # fmt: skip
        for lot in selected:
            counters[ws.title] += 1
            n = counters[ws.title]
            sources = Counter()
            natures = Counter()
            for need in lot.needs:
                line = lines[need.line_id]
                sources[line.funding_source] += need.amount
                if line.budget_line:
                    natures[line.budget_line] += need.amount
            dates = {s.label: s.end.strftime("%d-%m-%Y") for s in lot.steps}
            origin = "; ".join(activities.get(a, a) for a in lot.activity_ids)
            values = [
                n,
                f"{lot.market_type.value}_{lot.structure_id[:8]}_{n:03d}",
                f"{lot.category.replace('_', ' ').capitalize()} ({len(lot.needs)} besoin(s))",
                lot.market_type.value,
                lot.procedure.code,
                lot.amount_ht,
                " ".join(s for s, _ in sources.most_common()),
                " ".join(nat for nat, _ in natures.most_common()),
                lot.control_body + (" — publication UEMOA" if lot.community_publication else ""),
                "Annuel",
                *[dates.get(label, "") for label in step_labels],
                origin,
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row=4 + n, column=col, value=value)
    return _save(wb)
